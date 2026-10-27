# -*- coding: utf-8 -*-
from docx import Document
from pathlib import Path

p = Path(r"e:\CODE\LEARNING\DUOC - INGENIERIA EN INFORMATICA\Seg. y Calidad del Software\cuidarte-mas\docs\EV1 Entrega\EV1_Informe_Plan_Pruebas_Cuidarte_Mas.docx")
d = Document(str(p))
out = Path(r"e:\CODE\LEARNING\DUOC - INGENIERIA EN INFORMATICA\Seg. y Calidad del Software\cuidarte-mas\docs\EV1 Entrega\_sk_cp_dump.txt")
lines = []
# paragraphs around intro end
for i in range(28, 35):
    lines.append(f"P{i}: {d.paragraphs[i].text[:200]}")
lines.append("\n=== CP TABLES T13-T18 ===")
for ti in range(13, 19):
    t = d.tables[ti]
    lines.append(f"\n--- TABLE {ti} ---")
    for ri, row in enumerate(t.rows):
        # unique cells (merged often duplicate)
        seen = []
        vals = []
        for c in row.cells:
            txt = (c.text or "").replace("\n", " | ").strip()
            if txt in seen and len(row.cells) > 1:
                continue
            seen.append(txt)
            vals.append(txt[:200])
        lines.append(f"R{ri}: " + " || ".join(vals))
out.write_text("\n".join(lines), encoding="utf-8")
print("done")
