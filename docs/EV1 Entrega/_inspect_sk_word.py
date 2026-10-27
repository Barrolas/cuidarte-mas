# -*- coding: utf-8 -*-
from docx import Document
from pathlib import Path

p = Path(r"e:\CODE\LEARNING\DUOC - INGENIERIA EN INFORMATICA\Seg. y Calidad del Software\cuidarte-mas\docs\EV1 Entrega\EV1_Informe_Plan_Pruebas_Cuidarte_Mas.docx")
d = Document(str(p))
out = Path(r"e:\CODE\LEARNING\DUOC - INGENIERIA EN INFORMATICA\Seg. y Calidad del Software\cuidarte-mas\docs\EV1 Entrega\_sk_outline.txt")
lines = []
for i, para in enumerate(d.paragraphs):
    t = para.text.strip()
    if not t:
        continue
    style = para.style.name if para.style else ""
    lines.append(f"{i}|{style}|{t}")
lines.append("\n=== TABLES ===")
for ti, table in enumerate(d.tables):
    preview = []
    for row in table.rows[:4]:
        cells = [(c.text or "").replace("\n", " / ")[:80] for c in row.cells]
        preview.append(" || ".join(cells[:5]))
    lines.append(f"T{ti} ({len(table.rows)}x{len(table.columns)}):\n  " + "\n  ".join(preview))
# keyword presence
blob = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
keys = [
    "/autenticacion", "delayMiddleware", "JMeter", "HALL-", "4444", "3333",
    "Postman", "OWASP", "medico.qa", "PruebaSegura", "NFR-PERF", "CP-01", "CP-05",
    "1.2", "1.3", "1.4", "1.5", "propósito", "Propósito", "objetivos", "Alcance",
    "Swagger", "k6", "Lighthouse", "JWT", "SQLi",
]
lines.append("\n=== KEYWORDS ===")
for k in keys:
    lines.append(f"{k}: {k in blob or k.lower() in blob.lower()}")
out.write_text("\n".join(lines), encoding="utf-8")
print("ok", len(lines))
