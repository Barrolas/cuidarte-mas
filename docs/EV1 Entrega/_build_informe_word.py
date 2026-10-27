# -*- coding: utf-8 -*-
"""Genera el Word AVA del plan de pruebas Cuidarte+ desde el Markdown unificado."""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "Informe_Plan_Pruebas_CuidartePlus.md"
OUT_PATH = ROOT / "EV1_Informe_Plan_Pruebas_Cuidarte_Mas.docx"
BACKUP_PATH = ROOT / "EV1_Informe_Plan_Pruebas_Cuidarte_Mas_BACKUP_sk.docx"

INTEGRANTES = [
    "Ari Araya",
    "Giannina Guerrero",
    "Skarlett Tropan",
    "Nicolás Barra",
]


def set_run_font(run, size_pt: float, bold: bool = False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_paragraph(p, size_pt: float = 11, bold: bool = False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    for run in p.runs:
        set_run_font(run, size_pt, bold=bold)


def add_styled_para(doc, text: str, size_pt=11, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
    p = doc.add_paragraph()
    # Inline markdown **bold**
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size_pt, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, size_pt, bold=False)
        else:
            run = p.add_run(part)
            set_run_font(run, size_pt, bold=bold)
    style_paragraph(p, size_pt=size_pt, bold=False, align=align, space_after=space_after)
    return p


def add_heading_custom(doc, text: str, level: int):
    p = doc.add_paragraph()
    run = p.add_run(text)
    size = 12 if level <= 2 else 11
    set_run_font(run, size, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if level > 1 else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(6)
    return p


def clean_cell(text: str) -> str:
    t = text.strip()
    t = t.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    t = re.sub(r"\*\*([^*]+)\*\*", r"\1", t)
    t = re.sub(r"`([^`]+)`", r"\1", t)
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    return t


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(clean_cell(h))
        set_run_font(run, 11, bold=True)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    for r_idx, row in enumerate(rows):
        for c_idx in range(len(headers)):
            val = row[c_idx] if c_idx < len(row) else ""
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(clean_cell(val))
            set_run_font(run, 11, bold=False)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    doc.add_paragraph()


def parse_table(lines, start_idx):
    """Parse markdown table starting at start_idx. Returns (headers, rows, next_idx)."""
    headers = [c.strip() for c in lines[start_idx].strip().strip("|").split("|")]
    i = start_idx + 1
    if i < len(lines) and re.match(r"^\|?\s*-+", lines[i]):
        i += 1
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        cols = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(cols)
        i += 1
    return headers, rows, i


def strip_md_links(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text


def build_cover(doc: Document):
    for _ in range(2):
        doc.add_paragraph()

    p = add_styled_para(doc, "Cuidarte+", size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    for run in p.runs:
        set_run_font(run, 14, bold=True)

    add_styled_para(
        doc,
        "Informe de seguridad y calidad en el desarrollo de software",
        size_pt=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_styled_para(
        doc,
        "ISY1102 — Seguridad y Calidad en el Desarrollo de Software",
        size_pt=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_styled_para(doc, "Ingeniería en Informática", size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_para(doc, "Sección: 001V", size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_para(doc, "Docente: Jose Sergio Collio Huennun", size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_para(doc, "Fecha: septiembre 2026", size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_styled_para(
        doc,
        "Evaluación Parcial Nº 1 — Diseñando un plan de pruebas seguro, legal y normativo",
        size_pt=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )

    add_styled_para(doc, "Integrantes", size_pt=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for name in INTEGRANTES:
        add_styled_para(doc, name, size_pt=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

    add_styled_para(
        doc,
        "Equipo de 4 integrantes autorizado por el docente.",
        size_pt=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )
    doc.add_page_break()


def convert_md_body(doc: Document, md_text: str):
    # Skip YAML-less front; start from Índice or first ## after integrantes
    lines = md_text.replace("\r\n", "\n").split("\n")

    # Find "## Índice"
    start = 0
    for i, line in enumerate(lines):
        if line.strip().startswith("## Índice") or line.strip().startswith("## Indice"):
            start = i
            break
        if line.strip().startswith("## 1."):
            start = i
            break

    i = start
    while i < len(lines):
        line = lines[i]
        raw = line.rstrip()
        s = raw.strip()

        if not s:
            i += 1
            continue

        if s == "---":
            i += 1
            continue

        # Skip internal meta annex if desired — keep it as anexo
        if s.startswith("## "):
            title = s[3:].strip()
            add_heading_custom(doc, title, 1)
            i += 1
            continue

        if s.startswith("### "):
            title = s[4:].strip()
            add_heading_custom(doc, title, 2)
            i += 1
            continue

        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*-+", lines[i + 1].strip()):
            headers, rows, i = parse_table(lines, i)
            add_table(doc, headers, rows)
            continue

        if s.startswith("> "):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q = lines[i].strip()
                if q.startswith("> "):
                    q = q[2:]
                elif q == ">":
                    q = ""
                quote_lines.append(q)
                i += 1
            text = strip_md_links(" ".join(quote_lines))
            add_styled_para(doc, text, size_pt=11)
            continue

        if s.startswith("- ") or s.startswith("* "):
            while i < len(lines) and (lines[i].strip().startswith("- ") or lines[i].strip().startswith("* ")):
                item = strip_md_links(lines[i].strip()[2:])
                p = add_styled_para(doc, item, size_pt=11)
                p.paragraph_format.left_indent = Cm(0.75)
                i += 1
            continue

        if re.match(r"^\d+\.\s", s):
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                item = strip_md_links(lines[i].strip())
                p = add_styled_para(doc, item, size_pt=11)
                p.paragraph_format.left_indent = Cm(0.75)
                i += 1
            continue

        # Accumulated paragraph (continue until blank / special)
        buf = [strip_md_links(s)]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith(">") or nxt.startswith("- ") or nxt.startswith("* ") or nxt == "---" or re.match(r"^\d+\.\s", nxt):
                break
            buf.append(strip_md_links(nxt))
            i += 1
        text = " ".join(buf)
        text = html.unescape(text)
        add_styled_para(doc, text, size_pt=11)


def main():
    md = MD_PATH.read_text(encoding="utf-8")

    # Backup original Word once if not already backed up
    if OUT_PATH.exists() and not BACKUP_PATH.exists():
        BACKUP_PATH.write_bytes(OUT_PATH.read_bytes())

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    build_cover(doc)
    convert_md_body(doc, md)

    # Anexos pointer page
    add_heading_custom(doc, "Anexos de evidencia (carpeta Anexos/)", 1)
    add_styled_para(
        doc,
        "Los siguientes anexos acompañan este informe y se entregan junto al paquete EV1: "
        "Informe OWASP ZAP (FE/API/resumen), Informe de hallazgos HALL-01 a HALL-05 (MD/PDF), "
        "plan y resultados JMeter CP-05 (CuidartePlus_CP05.jmx, results_cp05.jtl, html_report e Informe_JMeter_CP05.md).",
    )

    doc.save(str(OUT_PATH))
    print(f"OK -> {OUT_PATH}")
    print(f"Bytes: {OUT_PATH.stat().st_size}")


if __name__ == "__main__":
    main()
