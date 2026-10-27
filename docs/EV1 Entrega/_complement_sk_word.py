# -*- coding: utf-8 -*-
"""
Complementa el Word de Skarlett con la parte técnica de Nicolás (cards),
sin reescribir el documento ni cambiar su formato base.
"""
from __future__ import annotations

import copy
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(
    r"e:\CODE\LEARNING\DUOC - INGENIERIA EN INFORMATICA\Seg. y Calidad del Software\cuidarte-mas\docs\EV1 Entrega"
)
BACKUP = ROOT / "EV1_Informe_Plan_Pruebas_Cuidarte_Mas_BACKUP_sk.docx"
TARGET = ROOT / "EV1_Informe_Plan_Pruebas_Cuidarte_Mas.docx"


def _apply_run_style(run, src_paragraph: Paragraph, bold: bool = False):
    src_runs = src_paragraph.runs
    if src_runs:
        src = src_runs[0]
        if src.font.name:
            run.font.name = src.font.name
            run._element.rPr.rFonts.set(qn("w:eastAsia"), src.font.name)
        if src.font.size:
            run.font.size = src.font.size
    if bold:
        run.bold = True


def clone_paragraph_after(paragraph: Paragraph, text: str, *, as_heading2: bool = False) -> Paragraph:
    new_p = copy.deepcopy(paragraph._p)
    for child in list(new_p):
        if child.tag == qn("w:r"):
            new_p.remove(child)
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if as_heading2:
        try:
            new_para.style = "Heading 2"
        except KeyError:
            pass
    run = new_para.add_run(text)
    _apply_run_style(run, paragraph, bold=as_heading2)
    return new_para


def insert_forward(anchor: Paragraph, items: list[tuple[str, bool]]) -> Paragraph:
    current = anchor
    for text, is_h2 in items:
        current = clone_paragraph_after(current, text, as_heading2=is_h2)
    return current


def insert_before(anchor: Paragraph, items: list[tuple[str, bool]]) -> None:
    for text, is_h2 in reversed(items):
        new_p = copy.deepcopy(anchor._p)
        for child in list(new_p):
            if child.tag == qn("w:r"):
                new_p.remove(child)
        anchor._p.addprevious(new_p)
        new_para = Paragraph(new_p, anchor._parent)
        if is_h2:
            try:
                new_para.style = "Heading 2"
            except KeyError:
                pass
        run = new_para.add_run(text)
        _apply_run_style(run, anchor, bold=is_h2)


def find_para(doc: Document, predicate):
    for i, p in enumerate(doc.paragraphs):
        if predicate(p.text.strip()):
            return i, p
    return None, None


def _append_to_field(table, label_pred, addition: str) -> bool:
    for ri, row in enumerate(table.rows):
        label = (row.cells[0].text or "").strip().upper()
        if not label_pred(label):
            continue
        content_row = table.rows[ri + 1]
        base = (content_row.cells[0].text or "").strip()
        if addition.strip() in base:
            return False
        new_text = f"{base} {addition.strip()}"
        for cell in content_row.cells:
            cell.text = new_text
        return True
    return False


def patch_cp_pasos(table, addition: str) -> bool:
    return _append_to_field(table, lambda l: l.startswith("PASOS"), addition)


def patch_cp_datos(table, addition: str) -> bool:
    return _append_to_field(table, lambda l: "DATOS" in l, addition)


def patch_cp_pre(table, addition: str) -> bool:
    return _append_to_field(table, lambda l: l.startswith("PRECOND"), addition)


def main():
    if not BACKUP.exists():
        raise SystemExit(f"Falta backup: {BACKUP}")
    shutil.copy2(BACKUP, TARGET)
    doc = Document(str(TARGET))

    # 1) Intro 1.2–1.5 (faltaba en Word Skarlett; requerido por plantilla)
    _, p_intro = find_para(doc, lambda t: t.startswith("Dado que el sistema trata datos"))
    if p_intro is None:
        raise SystemExit("No se encontró ancla de introducción")
    if find_para(doc, lambda t: t.startswith("1.2 Propósito"))[1] is None:
        insert_forward(
            p_intro,
            [
                ("1.2 Propósito del plan de pruebas", True),
                (
                    "Diseñar una estrategia integral que asegure, antes de producción, que Cuidarte+ cumple los requerimientos funcionales del ERS (RF-1 a RF-5), satisface atributos no funcionales (seguridad, rendimiento, usabilidad, disponibilidad y compatibilidad) y se alinea con el marco legal chileno de datos personales, ciberseguridad, documentos electrónicos y accesibilidad. El plan adopta el enfoque QA + Seguridad + Cumplimiento, evitando tratar la norma como una fase aislada al final.",
                    False,
                ),
                ("1.3 Objetivos", True),
                (
                    "Definir criterios de calidad, seguridad y cumplimiento trazables a códigos del ERS; clasificar y justificar los tipos de prueba aplicables; establecer estrategia, herramientas, criterios de entrada y salida, recursos y cronograma; diseñar al menos cinco casos representativos (CP-01 a CP-06); y evaluar cobertura, brechas y coherencia frente al ERS y al marco normativo (IE5).",
                    False,
                ),
                ("1.4 Alcance", True),
                (
                    "Incluye autenticación y sesión (RF-1), roles y reglas RB-1/RB-2/RB-3, exámenes y documentos (RF-3, RF-4), auditoría (RF-5) y atributos NFR-SEG, NFR-PERF, NFR-USAB y NFR-COMPAT. Quedan fuera de esta EV1 el pentest certificado, la certificación WCAG formal, la auditoría legal externa y la medición prolongada de disponibilidad 99,5%.",
                    False,
                ),
                ("1.5 Estructura del informe", True),
                (
                    "El documento sigue el Ítem I del encargo y la plantilla institucional: introducción; criterios y normativa; plan de pruebas (estrategia, tipos, aceptación, herramientas y recursos); diseño de casos; análisis de cobertura; conclusiones y referencias. Al final de la sección de casos se incorpora una nota técnica de ejecución exploratoria (API autenticada, OWASP ZAP y JMeter) como evidencia complementaria.",
                    False,
                ),
            ],
        )

    # 2) Endpoints / puertos (Card 6)
    _, justif = find_para(doc, lambda t: t.startswith("La justificación es doble"))
    if justif is not None and find_para(doc, lambda t: "Endpoints de referencia del backend" in t)[1] is None:
        insert_forward(
            justif,
            [
                ("Nota técnica de implementación (ambiente local)", True),
                (
                    "Endpoints de referencia del backend Cuidarte+ utilizados en el diseño y ejecución de casos: POST /autenticacion/login y POST /autenticacion/registro; /usuarios; /pacientes; /examenes; /documentos; /auditoria; documentación OpenAPI en /docs. En el despliegue Docker local del repositorio, el frontend opera en el puerto 3333 y la API en el puerto 4444 (Swagger: http://localhost:4444/docs).",
                    False,
                ),
            ],
        )

    # 3) Complemento en tablas CP (Cards 8–9)
    patches = {
        13: {
            "pasos": "Rutas reales de referencia: POST /autenticacion/login (API http://localhost:4444).",
        },
        14: {
            "pasos": "Rutas reales: POST /autenticacion/login; GET /pacientes; GET /tipo-examen; POST /examenes; verificación con GET /examenes/:id o GET /examenes/paciente/:paciente_id. Adjunto vía /documentos cuando corresponda.",
        },
        15: {
            "pasos": "Rutas reales: DELETE /examenes/:id (o DELETE /usuarios/:id según rol); consulta GET /auditoria.",
        },
        16: {
            "pasos": "Rutas reales: con token de paciente invocar GET /usuarios y/o DELETE /usuarios/:id (esperado 401/403); control positivo con token médico sobre GET /examenes/paciente/:id. Evidencia exploratoria sep-2026: paciente→GET /usuarios→403; login SQLi→401 sin JWT.",
        },
        17: {
            "pre": "Aviso técnico del repositorio: varias rutas (/examenes, /pacientes, /documentos) usan delayMiddleware(5000) en desarrollo; la medición formal de NFR-PERF-1 debe ejecutarse sin ese delay, o documentarse explícitamente como limitación del ambiente.",
            "pasos": "Artefacto ejecutable del equipo: Anexos/jmeter/CuidartePlus_CP05.jmx (TG-A 20 VUs; TG-B 50 VUs smoke de concurrencia).",
            "datos": "Resultado ejecutado con delay artificial activo: 220 samples, 0% error; login p95 ≈ 14 ms (cumple); GET /pacientes y GET /examenes p95 ≈ 5008–5011 ms (no cumple NFR-PERF-1). Ver Anexos/jmeter/Informe_JMeter_CP05.md.",
        },
        18: {
            "pasos": "Ambiente de referencia: frontend http://localhost:3333.",
        },
    }
    for ti, conf in patches.items():
        table = doc.tables[ti]
        if conf.get("pasos"):
            patch_cp_pasos(table, conf["pasos"])
        if conf.get("pre"):
            patch_cp_pre(table, conf["pre"])
        if conf.get("datos"):
            patch_cp_datos(table, conf["datos"])

    # 4) Hallazgos / evidencia (antes de §5)
    _, sec5 = find_para(doc, lambda t: t.startswith("5. Análisis de cobertura"))
    if sec5 is None:
        raise SystemExit("No se encontró sección 5")
    if find_para(doc, lambda t: t.startswith("4.1 Nota de ejecución exploratoria"))[1] is None:
        insert_before(
            sec5,
            [
                ("4.1 Nota de ejecución exploratoria (complemento técnico)", True),
                (
                    "Además del diseño de CP-01 a CP-06, se ejecutó una batería exploratoria sobre el ambiente Docker local (API autenticada, baseline OWASP ZAP sin sesión y JMeter) para sustentar los casos no funcionales y documentar defectos reales sin corregir código en esta entrega. Resumen de hallazgos:",
                    False,
                ),
                (
                    "HALL-01 (alta, funcional/usabilidad): el control «Ver Documento» en la UI de edición de examen no visualiza el adjunto de forma usable (se reporta apertura del cargador de archivos); la API GET /documentos/:id responde 200. Relación: RF-4 / NFR-USAB-2 · CP-02/CP-06.",
                    False,
                ),
                (
                    "HALL-02 (crítica, seguridad): contraseñas en texto plano (comparación e inserción sin bcrypt), en incumplimiento de NFR-SEG-2. Relación: CP-01/CP-04.",
                    False,
                ),
                (
                    "HALL-03 (alta, seguridad): JWT_SECRET con valor por defecto «inseguro», debilitando NFR-SEG-9.",
                    False,
                ),
                (
                    "HALL-04 (alta, rendimiento): delayMiddleware(5000) falsea NFR-PERF-1; JMeter midió p95 ≈ 5 s en GET de pacientes/exámenes. Relación: CP-05.",
                    False,
                ),
                (
                    "HALL-05 (media, hardening): baseline ZAP sin autenticación reportó ausencia/debilidad de headers (CSP, anti-clickjacking, X-Content-Type-Options, X-Powered-By, CORS amplio), alineado a NFR-SEG-5 · CP-04. El baseline sin sesión no alcanza flujos clínicos con JWT; por ello no evidencia por sí solo los defectos HALL-01 a HALL-04.",
                    False,
                ),
                (
                    "Evidencia anexa: Informe_Hallazgos_CuidartePlus (MD/PDF), reportes OWASP ZAP en Anexos/ e Informe_JMeter_CP05.md.",
                    False,
                ),
            ],
        )

    # 5) Recomendación delay/JMeter
    _, reco = find_para(doc, lambda t: t.startswith("Ampliar la suite con casos de cierre"))
    if reco is not None and find_para(doc, lambda t: t.startswith("Medir el rendimiento con delayMiddleware"))[1] is None:
        insert_forward(
            reco,
            [
                (
                    "Medir el rendimiento con delayMiddleware desactivado y reutilizar el plan JMeter de Anexos/jmeter/CuidartePlus_CP05.jmx para verificar NFR-PERF-1 en condiciones reales de QA.",
                    False,
                ),
            ],
        )

    # 6) Índice
    _, idx_line = find_para(doc, lambda t: t == "4. Diseño de casos de prueba")
    if idx_line is not None and find_para(doc, lambda t: t.startswith("4.1 Nota de ejecución"))[1] is None:
        insert_forward(idx_line, [("4.1 Nota de ejecución exploratoria (complemento técnico)", False)])

    doc.save(str(TARGET))
    print(f"OK complementado: {TARGET}")
    print(f"Bytes: {TARGET.stat().st_size}")


if __name__ == "__main__":
    main()
